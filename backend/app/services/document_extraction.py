from __future__ import annotations

from io import BytesIO
from pathlib import Path


def extract_text(filename: str, content: bytes) -> str:
    """Extract readable text from common road-maintenance office documents.

    Unsupported formats return an empty string rather than inventing content.
    """
    suffix = Path(filename).suffix.lower()

    if suffix in {".txt", ".csv"}:
        return content.decode("utf-8", errors="replace").strip()

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
        return "\n\n".join(page for page in pages if page).strip()

    if suffix == ".docx":
        from docx import Document

        document = Document(BytesIO(content))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        return "\n".join(paragraphs).strip()

    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            lines.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    lines.append(" | ".join(values))
        return "\n".join(lines).strip()

    return ""
